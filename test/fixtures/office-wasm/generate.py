# Synthetic fixture provenance only; NOT run by setup, qualification or CI.
# Original generator environment: python-docx==1.2.0, python-pptx==1.0.2,
# openpyxl==3.1.5, Pillow==12.2.0. No Python dependencies are needed by the lane.
# If manually regenerating, review resulting binary changes before accepting.
import importlib.util
import json
import mimetypes
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

# Use Python's built-in MIME registry, never host Apache configuration.
mimetypes.knownfiles = []
mimetypes.init(files=[])

root = Path(__file__).resolve().parent
image = Image.new("RGB", (400, 140), "white")
draw = ImageDraw.Draw(image)
draw.rectangle((10, 10, 390, 130), fill="#0f766e")
draw.rectangle((35, 40, 90, 110), fill="#facc15")
draw.rectangle((115, 25, 170, 110), fill="#60a5fa")
image.save(root / "fixture-image.png")
document = Document()
document.styles["Normal"].font.name = "Liberation Sans"
document.styles["Normal"].font.size = Pt(11)
document.sections[0].header.paragraphs[0].text = "COWORK WASM HEADER"
document.sections[0].footer.paragraphs[0].text = "COWORK WASM FOOTER"
document.add_heading("COWORK DOCX WASM PROOF", 0)
document.add_paragraph("Offline sandbox conversion marker: DOCX-7319.")
table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ["Item", "Units", "Amount"]):
    cell.text = text
for row in [("Alpha", "3", "120"), ("Beta", "4", "240")]:
    for cell, text in zip(table.add_row().cells, row):
        cell.text = text
document.add_picture(str(root / "fixture-image.png"), width=Inches(3))
document.add_page_break()
document.add_heading("SECOND PAGE DOCX-8420", 1)
document.add_paragraph("Explicit page break and nonblank second page.")
document.save(root / "fixture.docx")
created = ["fixture.docx"]
if importlib.util.find_spec("pptx"):
    from pptx import Presentation
    from pptx.util import Inches as PptInches
    presentation = Presentation()
    for marker in ["PPTX FIRST SLIDE 7319", "PPTX SECOND SLIDE 8420"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = marker
        slide.shapes.add_picture(str(root / "fixture-image.png"), PptInches(1), PptInches(2), width=PptInches(4))
        box = slide.shapes.add_textbox(PptInches(1), PptInches(4), PptInches(7), PptInches(1))
        box.text_frame.text = "Offline sandbox presentation proof."
    presentation.save(root / "fixture.pptx")
    created.append("fixture.pptx")
if importlib.util.find_spec("openpyxl"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.chart import BarChart, Reference
    book = Workbook()
    sheet = book.active
    sheet.title = "Proof"
    for row in [["XLSX SANDBOX PROOF", "Amount"], ["Alpha", 120], ["Beta", 240], ["TOTAL", "=SUM(B2:B3)"]]:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 18
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    chart = BarChart()
    chart.title = "Amounts"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    sheet.print_area = "A1:M18"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    book.save(root / "fixture.xlsx")
    created.append("fixture.xlsx")
root.joinpath("fixtures.json").write_text(json.dumps(created, indent=2))
print(json.dumps({"created": created}))
